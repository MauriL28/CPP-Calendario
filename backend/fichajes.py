"""Carga de fichajes: un evento por fila del Excel y las horas de la jornada.

Al terminar, cruza cada jornada con el turno efectivo y rellena coincide.
Horas en NUMERIC(8,2): segundos / 3600, redondeo half-up a 2 decimales.
Un tramo que cruza medianoche usa las dos fechas; no se suman 24 h a ciegas.
"""

import sys
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from decimal import Decimal, ROUND_HALF_UP

import openpyxl

from db import cursor
from rutas.turnos import _agrupar_versiones, _franja, _habituales, resolver_turno

TIPOS = ("Llegada", "Salida", "Inicio Pausa", "Final Pausa")
MAX_HUECO_JORNADA_H = 4
MAX_JORNADA_H = 16
MARGEN_DESCUADRE_MINUTOS = 20
COLUMNAS = (
    "Nº pers.",
    "Nombre del empleado o candidat",
    "Texto CHT",
    "Fe.lóg.",
    "Hora",
)


def _sap(valor):
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    if isinstance(valor, int):
        return str(valor)
    if valor is None:
        return ""
    return str(valor).strip()


def _tipo(valor):
    if valor is None:
        return None
    texto = str(valor).strip()
    return texto or None


def _fecha(valor):
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    raise ValueError(f"Fe.lóg. no es una fecha: {valor!r}")


def _hora(valor):
    if isinstance(valor, datetime):
        return valor.time()
    if isinstance(valor, time):
        return valor
    raise ValueError(f"Hora ausente o no válida: {valor!r}")


def leer_excel(ruta):
    """Devuelve (filas_leidas, filas_saltadas, eventos). No escribe el xlsx."""
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        if "Data" not in wb.sheetnames:
            raise ValueError("el libro no tiene la hoja Data")
        filas = wb["Data"].iter_rows(values_only=True)
        cabecera = next(filas, None)
        if cabecera is None:
            raise ValueError("la hoja Data está vacía")
        indice = {nombre: i for i, nombre in enumerate(cabecera)}
        faltan = [nombre for nombre in COLUMNAS if nombre not in indice]
        if faltan:
            raise ValueError("faltan columnas: " + ", ".join(faltan))
        leidas = 0
        saltadas = 0
        eventos = []
        for numero_fila, fila in enumerate(filas, start=2):
            leidas += 1
            tipo = _tipo(fila[indice["Texto CHT"]])
            fecha = _fecha(fila[indice["Fe.lóg."]])
            if tipo is None and fecha is None:
                saltadas += 1
                continue
            sap = _sap(fila[indice["Nº pers."]])
            if not sap:
                raise ValueError(f"fila {numero_fila}: Nº pers. vacío")
            eventos.append(
                {
                    "orden": numero_fila,
                    "numero_sap": sap,
                    "tipo": tipo,
                    "fecha": fecha,
                    "hora": _hora(fila[indice["Hora"]]),
                }
            )
        return leidas, saltadas, eventos
    finally:
        wb.close()


def _segundos(inicio, fin):
    def total(valor):
        entero = valor.hour * 3600 + valor.minute * 60 + valor.second
        return Decimal(entero) + Decimal(valor.microsecond) / Decimal(1000000)

    return total(fin) - total(inicio)


def _horas(segundos):
    return (segundos / Decimal(3600)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _delta_marca(inicio, fin):
    """Segundos entre dos eventos. Con fecha, la salida puede caer al día siguiente."""
    if inicio.get("fecha") is not None and fin.get("fecha") is not None:
        delta = datetime.combine(fin["fecha"], fin["hora"]) - datetime.combine(
            inicio["fecha"], inicio["hora"]
        )
        return (
            Decimal(delta.days) * Decimal(86400)
            + Decimal(delta.seconds)
            + Decimal(delta.microseconds) / Decimal(1000000)
        )
    return _segundos(inicio["hora"], fin["hora"])


def _incidencia(motivo):
    return {
        "incidencia": True,
        "motivo": motivo,
        "horas_trabajadas": None,
        "horas_pausa": None,
        "segundos_trabajo": None,
        "segundos_pausa": None,
    }


def calcular_dia(eventos):
    """Dos máquinas independientes. Empiezan cerradas. Eventos ya ordenados.

    Abrir abierto, cerrar cerrado, seguir abierto al final, tipo vacío o
    distinto de los cuatro, o un cierre anterior a su apertura: incidencia
    y las horas quedan vacías. Si los eventos traen fecha, el cierre del
    día siguiente cuenta el tiempo real.
    """
    trabajo_desde = None
    pausa_desde = None
    segundos_trabajo = Decimal(0)
    segundos_pausa = Decimal(0)
    for ev in eventos:
        tipo = ev["tipo"]
        hora = ev["hora"]
        if tipo not in TIPOS:
            return _incidencia(f"tipo no válido ({tipo or 'vacío'}) a las {hora}")
        if tipo == "Llegada":
            if trabajo_desde is not None:
                return _incidencia(
                    f"Llegada a las {hora} con la jornada ya abierta desde {trabajo_desde['hora']}"
                )
            trabajo_desde = ev
        elif tipo == "Salida":
            if trabajo_desde is None:
                return _incidencia(f"Salida a las {hora} sin Llegada")
            delta = _delta_marca(trabajo_desde, ev)
            if delta < 0:
                return _incidencia(
                    f"Salida a las {hora} anterior a la Llegada de las {trabajo_desde['hora']}"
                )
            segundos_trabajo += delta
            trabajo_desde = None
        elif tipo == "Inicio Pausa":
            if pausa_desde is not None:
                return _incidencia(
                    f"Inicio Pausa a las {hora} con la pausa ya abierta desde {pausa_desde['hora']}"
                )
            pausa_desde = ev
        else:
            if pausa_desde is None:
                return _incidencia(f"Final Pausa a las {hora} sin Inicio Pausa")
            delta = _delta_marca(pausa_desde, ev)
            if delta < 0:
                return _incidencia(
                    f"Final Pausa a las {hora} anterior al Inicio Pausa de las {pausa_desde['hora']}"
                )
            segundos_pausa += delta
            pausa_desde = None
    abiertos = []
    if trabajo_desde is not None:
        abiertos.append(f"Llegada a las {trabajo_desde['hora']} sin Salida")
    if pausa_desde is not None:
        abiertos.append(f"Inicio Pausa a las {pausa_desde['hora']} sin Final Pausa")
    if abiertos:
        return _incidencia("; ".join(abiertos))
    return {
        "incidencia": False,
        "motivo": None,
        "horas_trabajadas": _horas(segundos_trabajo),
        "horas_pausa": _horas(segundos_pausa),
        "segundos_trabajo": segundos_trabajo,
        "segundos_pausa": segundos_pausa,
    }


def _desde_medianoche(hora):
    entero = hora.hour * 3600 + hora.minute * 60 + hora.second
    return Decimal(entero) + Decimal(hora.microsecond) / Decimal(1000000)


def segundos_pausa_con_huecos(eventos):
    """Pausas fichadas más huecos Salida→Llegada de la misma lista, sin solapar.

    La lista es una jornada: un hueco que abre otra jornada no entra aquí.
    """
    if not eventos:
        return None
    ordenados = sorted(
        eventos,
        key=lambda item: (item.get("fecha") or date.min, item["hora"], item["orden"]),
    )

    def a_segundos(ev):
        if ev.get("fecha") is None or ordenados[0].get("fecha") is None:
            return _desde_medianoche(ev["hora"])
        delta = datetime.combine(ev["fecha"], ev["hora"]) - datetime.combine(
            ordenados[0]["fecha"], ordenados[0]["hora"]
        )
        return (
            Decimal(delta.days) * Decimal(86400)
            + Decimal(delta.seconds)
            + Decimal(delta.microseconds) / Decimal(1000000)
        )

    pausas = []
    huecos = []
    pausa_desde = None
    salida = None
    for ev in ordenados:
        tipo = ev["tipo"]
        if tipo == "Inicio Pausa" and pausa_desde is None:
            pausa_desde = ev
        elif tipo == "Final Pausa" and pausa_desde is not None:
            if a_segundos(ev) >= a_segundos(pausa_desde):
                pausas.append((pausa_desde, ev))
            pausa_desde = None
        elif tipo == "Salida":
            salida = ev
        elif tipo == "Llegada" and salida is not None:
            if a_segundos(ev) >= a_segundos(salida):
                huecos.append((salida, ev))
            salida = None
    total = sum((a_segundos(fin) - a_segundos(inicio) for inicio, fin in pausas), Decimal(0))
    for inicio, fin in huecos:
        desde = a_segundos(inicio)
        hasta = a_segundos(fin)
        cubierto = Decimal(0)
        for pausa_inicio, pausa_fin in pausas:
            cubierto += max(
                Decimal(0),
                min(hasta, a_segundos(pausa_fin)) - max(desde, a_segundos(pausa_inicio)),
            )
        total += max(Decimal(0), (hasta - desde) - cubierto)
    return total


def _momento(ev):
    return datetime.combine(ev["fecha"], ev["hora"])


def _partir_jornadas(lista):
    """Tramo = Llegada→Salida. Sigue en la jornada si el hueco es de 4 h o menos.

    La Salida casa con la Llegada abierta solo dentro de MAX_JORNADA_H.
    Una salida suelta, una pausa abierta o un tipo raro quedan en su propio grupo.
    """
    lista = sorted(lista, key=lambda ev: (ev["fecha"], ev["hora"], ev["orden"]))
    jornadas = []
    abierta = None
    llegada_dt = None
    salida_dt = None
    actual = None
    hueco = timedelta(hours=MAX_HUECO_JORNADA_H)
    tope = timedelta(hours=MAX_JORNADA_H)
    for ev in lista:
        momento = _momento(ev)
        tipo = ev["tipo"]
        if tipo == "Llegada":
            if abierta is not None:
                jornadas[abierta].append(ev)
                continue
            if salida_dt is not None and timedelta(0) <= (momento - salida_dt) <= hueco:
                jornadas[actual].append(ev)
                abierta = actual
                llegada_dt = momento
                salida_dt = None
                continue
            jornadas.append([ev])
            actual = len(jornadas) - 1
            abierta = actual
            llegada_dt = momento
            salida_dt = None
            continue
        if tipo == "Salida":
            if (
                abierta is not None
                and llegada_dt is not None
                and timedelta(0) <= (momento - llegada_dt) <= tope
            ):
                jornadas[abierta].append(ev)
                salida_dt = momento
                abierta = None
                llegada_dt = None
                continue
            if abierta is not None and llegada_dt is not None and (momento - llegada_dt) > tope:
                abierta = None
                llegada_dt = None
                salida_dt = None
                actual = None
            jornadas.append([ev])
            actual = None
            salida_dt = None
            continue
        if abierta is not None:
            jornadas[abierta].append(ev)
        elif actual is not None and salida_dt is not None:
            jornadas[actual].append(ev)
        else:
            jornadas.append([ev])
            actual = None
            salida_dt = None
    return jornadas


def _dia_de(persona, lista):
    calculo = calcular_dia(lista)
    llegada = next((ev for ev in lista if ev["tipo"] == "Llegada"), None)
    ancla = llegada or lista[0]
    return {
        "persona": persona,
        "usuario_id": ancla["usuario_id"],
        "numero_sap": ancla["numero_sap"],
        "fecha": ancla.get("fecha"),
        "eventos": lista,
        **calculo,
    }


def _rango_eventos(eventos):
    """Primer y último evento con fecha. None si no hay ninguno."""
    momentos = [_momento(ev) for ev in eventos if ev.get("fecha") is not None]
    if not momentos:
        return None
    return min(momentos), max(momentos)


def _llegadas_cerradas(dias):
    """Llegadas ya emparejadas con una Salida, por numero_sap."""
    cerradas = defaultdict(list)
    for dia in dias:
        trabajo = None
        for ev in dia["eventos"]:
            if ev["tipo"] == "Llegada":
                if trabajo is not None:
                    break
                trabajo = ev
            elif ev["tipo"] == "Salida":
                if trabajo is None or _momento(ev) < _momento(trabajo):
                    break
                cerradas[dia["numero_sap"]].append(_momento(trabajo))
                trabajo = None
    return cerradas


def _en_borde(dia, inicio, fin):
    """Huérfano al principio del extracto, o llegada abierta al final."""
    tope = timedelta(hours=MAX_JORNADA_H)
    eventos = dia["eventos"]
    llegadas = [ev for ev in eventos if ev["tipo"] == "Llegada"]
    salidas = [ev for ev in eventos if ev["tipo"] == "Salida"]
    if len(llegadas) == 1 and not salidas and _momento(llegadas[0]) >= fin - tope:
        return True
    if llegadas:
        return False
    if not eventos or not all(ev["tipo"] in ("Salida", "Inicio Pausa", "Final Pausa") for ev in eventos):
        return False
    return all(_momento(ev) <= inicio + tope for ev in eventos)


def _llegada_reabierta(eventos):
    """Hay una Llegada nueva mientras la anterior sigue sin Salida."""
    abierta = False
    for ev in eventos:
        if ev["tipo"] == "Llegada":
            if abierta:
                return True
            abierta = True
        elif ev["tipo"] == "Salida":
            abierta = False
    return False


def _motivo_incidencia(dia, cerradas):
    tope = timedelta(hours=MAX_JORNADA_H)
    eventos = dia["eventos"]
    llegadas = [ev for ev in eventos if ev["tipo"] == "Llegada"]
    salidas = [ev for ev in eventos if ev["tipo"] == "Salida"]
    if (dia.get("motivo") or "").startswith("tipo no válido"):
        return dia["motivo"]
    if _llegada_reabierta(eventos) or (llegadas and not salidas):
        return "Falta salida"
    if salidas and not llegadas:
        momento = _momento(salidas[0])
        hay_cerrada = any(
            timedelta(0) < momento - llegada <= tope
            for llegada in cerradas.get(dia["numero_sap"], [])
        )
        if hay_cerrada:
            return "Falta llegada tras salida"
        return "Falta entrada"
    return "Pausa sin cerrar"


def _marcar_fuera_de_rango(dia):
    dia.update(
        {
            "incidencia": False,
            "motivo": "fuera de rango",
            "horas_trabajadas": Decimal("0.00"),
            "horas_pausa": Decimal("0.00"),
            "segundos_trabajo": Decimal(0),
            "segundos_pausa": Decimal(0),
        }
    )


def _clasificar_motivos(dias, rango):
    if rango is None:
        return
    inicio, fin = rango
    cerradas = _llegadas_cerradas(dias)
    for dia in dias:
        if not dia["incidencia"]:
            continue
        if _en_borde(dia, inicio, fin):
            _marcar_fuera_de_rango(dia)
        else:
            dia["motivo"] = _motivo_incidencia(dia, cerradas)


def agrupar_dias(eventos, rango=None):
    """Una jornada por persona, en orden (fecha, hora, orden).

    La fecha de la jornada es la de su primera Llegada. Sin Llegada, la del
    primer evento. Persona: usuario_id si está, si no el numero_sap.
    rango es (primer evento, último evento) de fichaje_evento. Si no se pasa,
    se calcula con estos eventos.
    """
    por_persona = defaultdict(list)
    sin_fecha = defaultdict(list)
    for ev in eventos:
        persona = ev["usuario_id"] if ev["usuario_id"] else ev["numero_sap"]
        if ev.get("fecha") is None:
            sin_fecha[persona].append(ev)
        else:
            por_persona[persona].append(ev)
    dias = []
    for persona, lista in por_persona.items():
        for jornada in _partir_jornadas(lista):
            dias.append(_dia_de(persona, jornada))
    for persona, lista in sin_fecha.items():
        lista.sort(key=lambda ev: (ev["hora"], ev["orden"]))
        dias.append(_dia_de(persona, lista))
    if rango is None:
        rango = _rango_eventos(eventos)
    _clasificar_motivos(dias, rango)
    return dias


def _mismo_reloj(izquierda, derecha):
    def partes(valor):
        if isinstance(valor, datetime):
            valor = valor.time()
        return (valor.hour, valor.minute, valor.second, valor.microsecond)

    return partes(izquierda) == partes(derecha)


def _igual_evento(previo, ev):
    return (
        previo["numero_sap"] == ev["numero_sap"]
        and previo["tipo"] == ev["tipo"]
        and previo["fecha"] == ev["fecha"]
        and _mismo_reloj(previo["hora"], ev["hora"])
    )


def _insertar_eventos(cur, eventos):
    """Devuelve (nuevas, actualizadas).

    La fila que ya está, con el mismo orden y los mismos datos, no se borra
    ni se vuelve a insertar. Así la segunda carga del mismo libro no cuenta
    como filas nuevas.
    """
    cur.execute(
        """
        SELECT orden, numero_sap, tipo, fecha, hora
        FROM fichaje_evento
        """
    )
    previos = {fila["orden"]: fila for fila in cur.fetchall()}
    ordenes = [ev["orden"] for ev in eventos]
    if ordenes:
        cur.execute(
            "DELETE FROM fichaje_evento WHERE NOT (orden = ANY(%s))",
            (ordenes,),
        )
    else:
        cur.execute("DELETE FROM fichaje_evento")
    nuevos = []
    cambiados = []
    for ev in eventos:
        previo = previos.get(ev["orden"])
        if previo is None:
            nuevos.append(ev)
        elif not _igual_evento(previo, ev):
            cambiados.append(ev)
    if nuevos:
        cur.executemany(
            """
            INSERT INTO fichaje_evento (orden, numero_sap, tipo, fecha, hora)
            VALUES (%s, %s, %s, %s, %s)
            """,
            [
                (ev["orden"], ev["numero_sap"], ev["tipo"], ev["fecha"], ev["hora"])
                for ev in nuevos
            ],
        )
    if cambiados:
        cur.executemany(
            """
            UPDATE fichaje_evento
            SET numero_sap = %s, tipo = %s, fecha = %s, hora = %s, usuario_id = NULL
            WHERE orden = %s
            """,
            [
                (ev["numero_sap"], ev["tipo"], ev["fecha"], ev["hora"], ev["orden"])
                for ev in cambiados
            ],
        )
    cur.execute(
        """
        UPDATE fichaje_evento AS e
        SET usuario_id = m.id
        FROM (
            SELECT numero_sap, (array_agg(id))[1] AS id
            FROM usuario
            WHERE numero_sap IS NOT NULL
            GROUP BY numero_sap
            HAVING COUNT(*) = 1
        ) AS m
        WHERE e.numero_sap = m.numero_sap
        """
    )
    return len(nuevos), len(cambiados)


def _leer_eventos(cur):
    cur.execute(
        """
        SELECT orden, numero_sap, usuario_id, tipo, fecha, hora
        FROM fichaje_evento
        ORDER BY orden
        """
    )
    return list(cur.fetchall())


def _margen():
    return Decimal(MARGEN_DESCUADRE_MINUTOS) / Decimal(60)


def _coincide_con_turno(horas_trabajadas, turno):
    if turno is None:
        return "sin_plan"
    planificadas = turno["horas_planificadas"]
    if not isinstance(planificadas, Decimal):
        planificadas = Decimal(str(planificadas))
    if abs(horas_trabajadas - planificadas) <= _margen():
        return "ok"
    return "descuadre"


def _turnos_resueltos(cur, dias):
    """Días con usuario y sin incidencia. La clave es (usuario_id, fecha)."""
    utiles = [
        dia
        for dia in dias
        if dia["usuario_id"] is not None and dia["fecha"] is not None and not dia["incidencia"]
    ]
    if not utiles:
        return {}
    ids = list({dia["usuario_id"] for dia in utiles})
    fechas = [dia["fecha"] for dia in utiles]
    cur.execute(
        """
        SELECT usuario_id, fecha, ausencia::text AS ausencia,
               hora_inicio, hora_fin, horas_planificadas, horas_nocturnas
        FROM turno
        WHERE usuario_id = ANY(%s)
          AND fecha = ANY(%s)
        """,
        (ids, fechas),
    )
    turnos = {(fila["usuario_id"], fila["fecha"]): fila for fila in cur.fetchall()}
    cur.execute(
        """
        SELECT v.usuario_id, v.desde, d.dia_semana, d.horario_inicio, d.horario_fin
        FROM horario_version v
        LEFT JOIN horario_dia d
          ON d.usuario_id = v.usuario_id AND d.desde = v.desde
        WHERE v.usuario_id = ANY(%s)
          AND v.desde <= %s
        """,
        (ids, max(fechas)),
    )
    versiones = _agrupar_versiones(cur.fetchall())
    noche_inicio, noche_fin = _franja(cur)
    habituales = _habituales(cur, ids)
    resueltos = {}
    for dia in utiles:
        fila = turnos.get((dia["usuario_id"], dia["fecha"]))
        turno, _origen = resolver_turno(
            dia["fecha"],
            fila,
            versiones.get(dia["usuario_id"], []),
            noche_inicio,
            noche_fin,
            habituales.get(dia["usuario_id"], True),
        )
        resueltos[(dia["usuario_id"], dia["fecha"])] = (turno, _origen)
    return resueltos


def _horas_plan(turno):
    if turno is None:
        return None
    plan = turno["horas_planificadas"]
    if not isinstance(plan, Decimal):
        plan = Decimal(str(plan))
    return plan.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _cruce(dia, turnos):
    if dia.get("motivo") == "fuera de rango":
        return None, None
    if dia["usuario_id"] is None:
        return "sin_emparejar", None
    if dia["incidencia"]:
        return None, None
    turno, _origen = turnos[(dia["usuario_id"], dia["fecha"])]
    return _coincide_con_turno(dia["horas_trabajadas"], turno), _horas_plan(turno)


def _guardar_fichaje(cur, dia, coincide, planificadas):
    cur.execute(
        """
        INSERT INTO fichaje (
            usuario_id, numero_sap, fecha, horas_fichadas,
            horas_trabajadas, horas_pausa, horas_planificadas, incidencia, coincide
        )
        VALUES (%s, %s, %s, 0, %s, %s, %s, %s, %s)
        """,
        (
            dia["usuario_id"],
            dia["numero_sap"],
            dia["fecha"],
            dia["horas_trabajadas"],
            dia["horas_pausa"],
            planificadas,
            dia["incidencia"],
            coincide,
        ),
    )
    return True


def _insertar_fichajes(cur, dias):
    """Cada jornada es una fila. Se reescriben todas: la fecha ya no es el día natural."""
    turnos = _turnos_resueltos(cur, dias)
    cur.execute("DELETE FROM fichaje")
    insertados = 0
    for dia in dias:
        if dia["fecha"] is None:
            continue
        coincide, planificadas = _cruce(dia, turnos)
        if _guardar_fichaje(cur, dia, coincide, planificadas):
            insertados += 1
    return insertados


def _numeros_ambiguos(cur):
    cur.execute(
        """
        SELECT u.numero_sap
        FROM usuario AS u
        JOIN (SELECT DISTINCT numero_sap FROM fichaje_evento) AS e
          ON e.numero_sap = u.numero_sap
        GROUP BY u.numero_sap
        HAVING COUNT(*) > 1
        """
    )
    return {fila["numero_sap"] for fila in cur.fetchall()}


def cargar(ruta, registro=None):
    leidas, saltadas, eventos = leer_excel(ruta)
    with cursor() as cur:
        nuevas, actualizadas = _insertar_eventos(cur, eventos)
        guardados = _leer_eventos(cur)
        ambiguos = _numeros_ambiguos(cur)
        dias = agrupar_dias(guardados)
        fichajes = _insertar_fichajes(cur, dias)
        cur.execute(
            "SELECT COUNT(*) AS n FROM usuario WHERE numero_sap IS NOT NULL"
        )
        usuarios_con_sap = cur.fetchone()["n"]
        cur.execute("SELECT COUNT(*) AS n FROM fichaje_evento")
        filas_evento = cur.fetchone()["n"]
        cur.execute("SELECT COUNT(*) AS n FROM fichaje")
        filas_fichaje = cur.fetchone()["n"]
        con_usuario = [ev for ev in guardados if ev["usuario_id"] is not None]
        sin_coincidencia = [
            ev
            for ev in guardados
            if ev["usuario_id"] is None and ev["numero_sap"] not in ambiguos
        ]
        filas_ambiguas = [ev for ev in guardados if ev["numero_sap"] in ambiguos]
        dias_resueltos = [dia for dia in dias if dia["usuario_id"] is not None]
        incidencias = sum(1 for dia in dias if dia["incidencia"])
        sin_emparejar = len(sin_coincidencia) + len(filas_ambiguas)
        if registro is not None:
            cur.execute(
                """
                INSERT INTO fichaje_importacion (
                    usuario_id, nombre_archivo, filas_leidas,
                    filas_insertadas_o_actualizadas,
                    filas_sin_emparejar, dias_incidencia
                )
                VALUES (%s, %s, %s, %s, %s, %s)
                """,
                (
                    registro["usuario_id"],
                    registro["nombre_archivo"],
                    leidas,
                    nuevas + actualizadas,
                    sin_emparejar,
                    incidencias,
                ),
            )
    return {
        "filas_leidas": leidas,
        "filas_saltadas": saltadas,
        "filas_insertadas": filas_evento,
        "filas_nuevas": nuevas,
        "filas_actualizadas": actualizadas,
        "filas_con_usuario": len(con_usuario),
        "filas_sin_coincidencia": len(sin_coincidencia),
        "filas_ambiguas": len(filas_ambiguas),
        "sin_emparejar": sin_emparejar,
        "numeros_ambiguos": len(ambiguos),
        "personas_excel": len({ev["numero_sap"] for ev in guardados}),
        "personas_emparejadas": len({ev["numero_sap"] for ev in con_usuario}),
        "usuarios_con_sap": usuarios_con_sap,
        "dias": len(dias),
        "dias_resueltos": len(dias_resueltos),
        "incidencias": incidencias,
        "incidencias_resueltas": sum(1 for dia in dias_resueltos if dia["incidencia"]),
        "fichajes_insertados": fichajes,
        "fichajes_en_tabla": filas_fichaje,
        "dias_detalle": dias,
    }


def _texto_eventos(dia):
    return ", ".join(f"{ev['tipo']} {ev['hora']}" for ev in dia["eventos"])


def _imprimir_dia(titulo, dia):
    if dia is None:
        print(f"{titulo}: ninguno")
        return
    print(
        f"{titulo}: sap={dia['numero_sap']} fecha={dia['fecha']} "
        f"trabajadas={dia['horas_trabajadas']} pausa={dia['horas_pausa']} "
        f"incidencia={dia['incidencia']} motivo={dia['motivo']}"
    )
    print(f"  {_texto_eventos(dia)}")


def main():
    if len(sys.argv) != 2:
        print("uso: python fichajes.py RUTA.xlsx", file=sys.stderr)
        return 2
    comprobar_calcular_dia()
    resultado = cargar(sys.argv[1])
    print(f"filas_leidas={resultado['filas_leidas']}")
    print(f"filas_saltadas={resultado['filas_saltadas']}")
    print(f"filas_insertadas={resultado['filas_insertadas']}")
    print(f"filas_nuevas={resultado['filas_nuevas']}")
    print(f"filas_actualizadas={resultado['filas_actualizadas']}")
    print(f"filas_con_usuario={resultado['filas_con_usuario']}")
    print(f"filas_sin_coincidencia={resultado['filas_sin_coincidencia']}")
    print(f"filas_ambiguas={resultado['filas_ambiguas']}")
    print(f"numeros_ambiguos={resultado['numeros_ambiguos']}")
    print(f"personas_excel={resultado['personas_excel']}")
    print(f"personas_emparejadas={resultado['personas_emparejadas']}")
    print(f"usuarios_con_sap={resultado['usuarios_con_sap']}")
    print(f"dias={resultado['dias']}")
    print(f"dias_resueltos={resultado['dias_resueltos']}")
    print(f"incidencias={resultado['incidencias']}")
    print(f"incidencias_resueltas={resultado['incidencias_resueltas']}")
    print(f"fichajes_insertados={resultado['fichajes_insertados']}")
    print(f"fichajes_en_tabla={resultado['fichajes_en_tabla']}")
    dias = resultado["dias_detalle"]
    normal = next(
        (
            dia
            for dia in sorted(dias, key=lambda d: (d["numero_sap"], d["fecha"] or date.min))
            if not dia["incidencia"]
            and dia["horas_trabajadas"] > 0
            and dia["horas_pausa"] > 0
        ),
        None,
    )
    incidencia = next(
        (
            dia
            for dia in sorted(dias, key=lambda d: (d["numero_sap"], d["fecha"] or date.min))
            if dia["incidencia"]
        ),
        None,
    )
    _imprimir_dia("ejemplo_normal", normal)
    _imprimir_dia("ejemplo_incidencia", incidencia)
    return 0


def comprobar_calcular_dia():
    normal = calcular_dia(
        [
            {"tipo": "Llegada", "hora": time(5, 59, 40)},
            {"tipo": "Inicio Pausa", "hora": time(12, 43, 38)},
            {"tipo": "Final Pausa", "hora": time(13, 2, 53)},
            {"tipo": "Salida", "hora": time(13, 42, 46)},
        ]
    )
    if (
        normal["incidencia"]
        or normal["horas_trabajadas"] != Decimal("7.72")
        or normal["horas_pausa"] != Decimal("0.32")
    ):
        raise AssertionError(normal)
    noche = calcular_dia(
        [
            {"tipo": "Llegada", "hora": time(22, 0, 0), "fecha": date(2026, 9, 7)},
            {"tipo": "Salida", "hora": time(0, 30, 0), "fecha": date(2026, 9, 8)},
        ]
    )
    if noche["incidencia"] or noche["horas_trabajadas"] != Decimal("2.50"):
        raise AssertionError(noche)
    atras = calcular_dia(
        [
            {"tipo": "Llegada", "hora": time(22, 0, 0), "fecha": date(2026, 9, 7)},
            {"tipo": "Salida", "hora": time(0, 30, 0), "fecha": date(2026, 9, 7)},
        ]
    )
    if not atras["incidencia"] or atras["horas_trabajadas"] is not None:
        raise AssertionError(atras)
    cruzada = agrupar_dias(
        [
            {
                "usuario_id": "u",
                "numero_sap": "1",
                "orden": 1,
                "tipo": "Llegada",
                "fecha": date(2026, 9, 7),
                "hora": time(20, 21, 0),
            },
            {
                "usuario_id": "u",
                "numero_sap": "1",
                "orden": 2,
                "tipo": "Salida",
                "fecha": date(2026, 9, 8),
                "hora": time(4, 0, 0),
            },
            {
                "usuario_id": "u",
                "numero_sap": "1",
                "orden": 3,
                "tipo": "Llegada",
                "fecha": date(2026, 9, 8),
                "hora": time(20, 22, 0),
            },
        ]
    )
    if (
        len(cruzada) != 2
        or cruzada[0]["incidencia"]
        or cruzada[0]["fecha"] != date(2026, 9, 7)
        or cruzada[1]["incidencia"]
        or cruzada[1]["motivo"] != "fuera de rango"
        or cruzada[1]["fecha"] != date(2026, 9, 8)
    ):
        raise AssertionError(cruzada)
    abierta = calcular_dia([{"tipo": "Llegada", "hora": time(20, 21, 10)}])
    if not abierta["incidencia"] or "sin Salida" not in abierta["motivo"]:
        raise AssertionError(abierta)
    rara = calcular_dia([{"tipo": None, "hora": time(8, 0, 0)}])
    if not rara["incidencia"] or rara["horas_trabajadas"] is not None:
        raise AssertionError(rara)


if __name__ == "__main__":
    sys.exit(main())
