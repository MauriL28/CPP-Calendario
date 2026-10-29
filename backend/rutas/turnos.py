from datetime import datetime, timedelta
from zoneinfo import ZoneInfo
from calendar import monthrange
from io import BytesIO

import psycopg
from flask import Blueprint, jsonify, request, send_file
from flask_jwt_extended import get_jwt, get_jwt_identity, jwt_required
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from db import SinBaseDeDatos, cursor
from errores import error_bd
from mandos import departamento_permitido
from horas import hora_texto, horas_de_turno, parse_hora

bp = Blueprint("turnos", __name__)

AUSENCIAS = {"L", "D", "V", "B", "F", "P"}


def _lunes(texto, admite_tipo):
    try:
        desde = datetime.strptime(texto, "%Y-%m-%d").date()
    except (TypeError, ValueError) if admite_tipo else ValueError:
        return None, (jsonify(error="desde no es una fecha válida"), 400)
    if desde.weekday() != 0:
        return None, (jsonify(error="desde tiene que ser lunes"), 400)
    return desde, None


def _turno(fila):
    return {
        "ausencia": fila["ausencia"],
        "hora_inicio": hora_texto(fila["hora_inicio"]),
        "hora_fin": hora_texto(fila["hora_fin"]),
        "horas_planificadas": float(fila["horas_planificadas"]),
        "horas_nocturnas": float(fila["horas_nocturnas"]),
    }


def _franja(cur):
    cur.execute("SELECT noche_inicio, noche_fin FROM empresa WHERE id = 1")
    franja = cur.fetchone()
    if franja is None:
        return (
            datetime.strptime("22:00", "%H:%M").time(),
            datetime.strptime("06:00", "%H:%M").time(),
        )
    return franja["noche_inicio"], franja["noche_fin"]


def _agrupar_versiones(filas):
    grupos = {}
    for fila in filas:
        clave = (fila["usuario_id"], fila["desde"])
        version = grupos.get(clave)
        if version is None:
            version = {"desde": fila["desde"], "dias": {}}
            grupos[clave] = version
        if fila["dia_semana"] is not None:
            version["dias"][fila["dia_semana"]] = (fila["horario_inicio"], fila["horario_fin"])
    por_usuario = {}
    for (usuario_id, _), version in grupos.items():
        por_usuario.setdefault(usuario_id, []).append(version)
    for versiones in por_usuario.values():
        versiones.sort(key=lambda version: version["desde"])
    return por_usuario


def _horas_habituales(fecha, versiones):
    version = None
    for candidata in versiones:
        if candidata["desde"] <= fecha:
            version = candidata
        else:
            break
    if version is None or not version["dias"]:
        return None
    return version["dias"].get(fecha.weekday())


def _domingo_libre():
    return {
        "ausencia": "D",
        "hora_inicio": None,
        "hora_fin": None,
        "horas_planificadas": 0.0,
        "horas_nocturnas": 0.0,
    }, "habitual"


def resolver_turno(fecha, fila, versiones, noche_inicio, noche_fin, usa_horario_habitual=True):
    if fila is not None:
        return _turno(fila), "guardado"
    if not usa_horario_habitual:
        return None, "vacio"
    horas = _horas_habituales(fecha, versiones)
    if horas is None:
        if fecha.weekday() == 6:
            return _domingo_libre()
        return None, None
    inicio, fin = horas
    planificadas, nocturnas = horas_de_turno(inicio, fin, noche_inicio, noche_fin)
    return {
        "ausencia": None,
        "hora_inicio": hora_texto(inicio),
        "hora_fin": hora_texto(fin),
        "horas_planificadas": planificadas,
        "horas_nocturnas": nocturnas,
    }, "habitual"


def turno_efectivo(usuario, fecha):
    with cursor() as cur:
        cur.execute(
            """
            SELECT ausencia::text AS ausencia, hora_inicio, hora_fin,
                   horas_planificadas, horas_nocturnas
            FROM turno
            WHERE usuario_id = %s AND fecha = %s
            """,
            (usuario, fecha),
        )
        fila = cur.fetchone()
        if fila is not None:
            return resolver_turno(fecha, fila, [], None, None)[0]
        cur.execute(
            """
            SELECT v.usuario_id, v.desde, d.dia_semana, d.horario_inicio, d.horario_fin
            FROM horario_version v
            LEFT JOIN horario_dia d
              ON d.usuario_id = v.usuario_id AND d.desde = v.desde
            WHERE v.usuario_id = %s AND v.desde <= %s
            """,
            (usuario, fecha),
        )
        agrupadas = _agrupar_versiones(cur.fetchall())
        versiones = next(iter(agrupadas.values()), [])
        noche_inicio, noche_fin = _franja(cur)
        habitual = _habituales(cur, [usuario]).get(usuario, True)
    return resolver_turno(fecha, None, versiones, noche_inicio, noche_fin, habitual)[0]


def _material_rango(cur, ids, inicio, fin):
    turnos = {}
    versiones = {}
    if ids:
        cur.execute(
            """
            SELECT usuario_id, fecha, ausencia::text AS ausencia,
                   hora_inicio, hora_fin, horas_planificadas, horas_nocturnas
            FROM turno
            WHERE usuario_id = ANY(%s)
              AND fecha >= %s
              AND fecha <= %s
            """,
            (ids, inicio, fin),
        )
        for fila in cur.fetchall():
            turnos[(fila["usuario_id"], fila["fecha"])] = fila
        cur.execute(
            """
            SELECT v.usuario_id, v.desde, d.dia_semana, d.horario_inicio, d.horario_fin
            FROM horario_version v
            LEFT JOIN horario_dia d
              ON d.usuario_id = v.usuario_id AND d.desde = v.desde
            WHERE v.usuario_id = ANY(%s)
              AND v.desde <= %s
            """,
            (ids, fin),
        )
        versiones = _agrupar_versiones(cur.fetchall())
    noche_inicio, noche_fin = _franja(cur)
    return turnos, versiones, noche_inicio, noche_fin, _habituales(cur, ids)


def _habituales(cur, ids):
    if not ids:
        return {}
    cur.execute(
        """
        SELECT u.id, d.usa_horario_habitual
        FROM usuario u
        JOIN departamento d ON d.id = u.departamento_id
        WHERE u.id = ANY(%s)
        """,
        (list(ids),),
    )
    return {fila["id"]: fila["usa_horario_habitual"] for fila in cur.fetchall()}


def _material_semana(cur, ids, dias):
    return _material_rango(cur, ids, dias[0], dias[6])


def _festivos_rango(cur, inicio, fin):
    cur.execute(
        """
        SELECT fecha, nombre
        FROM festivo
        WHERE fecha >= %s AND fecha <= %s
        """,
        (inicio, fin),
    )
    return {fila["fecha"]: fila["nombre"] for fila in cur.fetchall()}


def _festivos_semana(cur, dias):
    return _festivos_rango(cur, dias[0], dias[6])


def _dias_resueltos(
    usuario_id, dias, turnos, versiones, noche_inicio, noche_fin, festivos, habituales=None
):
    suyas = versiones.get(usuario_id, [])
    habitual = True if habituales is None else habituales.get(usuario_id, True)
    resultado = []
    for dia in dias:
        turno, origen = resolver_turno(
            dia,
            turnos.get((usuario_id, dia)),
            suyas,
            noche_inicio,
            noche_fin,
            habitual,
        )
        entrada = {"fecha": dia.isoformat(), "turno": turno, "festivo": dia in festivos}
        if origen is not None:
            entrada["origen"] = origen
        if dia in festivos:
            entrada["nombre"] = festivos[dia]
        resultado.append(entrada)
    return resultado


def _semana(desde):
    return [desde + timedelta(days=i) for i in range(7)]


def _mes(anio_texto, mes_texto):
    if anio_texto in (None, "") or mes_texto in (None, ""):
        return None, (jsonify(error="Faltan datos"), 400)
    try:
        anio = int(anio_texto)
        mes = int(mes_texto)
        inicio = datetime(anio, mes, 1).date()
    except (TypeError, ValueError):
        return None, (jsonify(error="La fecha no es válida"), 400)
    ultimo = monthrange(anio, mes)[1]
    return [inicio + timedelta(days=i) for i in range(ultimo)], None


def _trabajadores_resueltos(
    trabajadores, dias, turnos, versiones, noche_inicio, noche_fin, festivos, habituales=None
):
    return [
        {
            "nombre": trabajador["nombre"],
            "login": trabajador["login"],
            "grupo": trabajador["grupo"],
            "dias": _dias_resueltos(
                trabajador["id"],
                dias,
                turnos,
                versiones,
                noche_inicio,
                noche_fin,
                festivos,
                habituales,
            ),
        }
        for trabajador in trabajadores
    ]


def _visible_en_semana(trabajador, dias, turnos):
    if trabajador["activo"]:
        return True
    return any((trabajador["id"], dia) in turnos for dia in dias)


def _departamento_lectura(cur):
    """Departamento del cuadrante. El propio si no llega departamento_id."""
    pedido = request.args.get("departamento_id")
    if pedido in (None, ""):
        cur.execute(
            """
            SELECT departamento_id
            FROM usuario
            WHERE id = %s AND rol = 'mando'
            """,
            (get_jwt_identity(),),
        )
        mando = cur.fetchone()
        if mando is None or mando["departamento_id"] is None:
            return None, (jsonify(error="No autorizado"), 403)
        return mando["departamento_id"], None
    if not departamento_permitido(cur, get_jwt_identity(), pedido):
        return None, (jsonify(error="No se ha encontrado el departamento"), 404)
    return pedido, None


@bp.get("/turnos")
@jwt_required()
def listar_turnos():
    if get_jwt().get("rol") != "mando":
        return jsonify(error="No autorizado"), 403
    desde, error = _lunes(request.args.get("desde", ""), False)
    if error:
        return error
    dias = _semana(desde)

    try:
        with cursor() as cur:
            departamento_id, error = _departamento_lectura(cur)
            if error:
                return error
            cur.execute(
                """
                SELECT u.id, u.nombre, u.login, u.grupo, u.activo, s.nombre AS seccion
                FROM usuario u
                LEFT JOIN seccion s ON s.id = u.seccion_id
                WHERE u.rol = 'trabajador' AND u.departamento_id = %s
                ORDER BY u.nombre
                """,
                (departamento_id,),
            )
            trabajadores = cur.fetchall()
            turnos, versiones, noche_inicio, noche_fin, habituales = _material_semana(
                cur, [trabajador["id"] for trabajador in trabajadores], dias
            )
            festivos = _festivos_semana(cur, dias)
            trabajadores = [
                trabajador
                for trabajador in trabajadores
                if _visible_en_semana(trabajador, dias, turnos)
            ]
    except SinBaseDeDatos:
        return jsonify(error="Base de datos no configurada"), 503
    except psycopg.Error as exc:
        return error_bd(exc)

    return jsonify(
        desde=desde.isoformat(),
        dias=[dia.isoformat() for dia in dias],
        trabajadores=[
            {
                "nombre": trabajador["nombre"],
                "login": trabajador["login"],
                "grupo": trabajador["grupo"],
                "seccion": trabajador["seccion"],
                "dias": _dias_resueltos(
                    trabajador["id"],
                    dias,
                    turnos,
                    versiones,
                    noche_inicio,
                    noche_fin,
                    festivos,
                    habituales,
                ),
            }
            for trabajador in trabajadores
        ],
    )


@bp.get("/turnos/mios")
@jwt_required()
def listar_turnos_mios():
    if get_jwt().get("rol") != "trabajador":
        return jsonify(error="No autorizado"), 403
    desde, error = _lunes(request.args.get("desde", ""), False)
    if error:
        return error
    dias = _semana(desde)

    try:
        with cursor() as cur:
            cur.execute(
                """
                SELECT id, nombre, login, grupo, activo
                FROM usuario
                WHERE id = %s AND rol = 'trabajador'
                """,
                (get_jwt_identity(),),
            )
            trabajador = cur.fetchone()
            if trabajador is None:
                return jsonify(error="No autorizado"), 403
            turnos, versiones, noche_inicio, noche_fin, habituales = _material_semana(
                cur, [trabajador["id"]], dias
            )
            festivos = _festivos_semana(cur, dias)
            if not _visible_en_semana(trabajador, dias, turnos):
                trabajador = None
    except SinBaseDeDatos:
        return jsonify(error="Base de datos no configurada"), 503
    except psycopg.Error as exc:
        return error_bd(exc)

    personas = []
    if trabajador is not None:
        personas.append(
            {
                "nombre": trabajador["nombre"],
                "login": trabajador["login"],
                "grupo": trabajador["grupo"],
                "dias": _dias_resueltos(
                    trabajador["id"],
                    dias,
                    turnos,
                    versiones,
                    noche_inicio,
                    noche_fin,
                    festivos,
                    habituales,
                ),
            }
        )
    return jsonify(
        desde=desde.isoformat(),
        dias=[dia.isoformat() for dia in dias],
        trabajadores=personas,
    )


def _turno_visible(turno, propio):
    if turno is None:
        return None
    ausencia = turno["ausencia"]
    if ausencia:
        if ausencia == "B" and not propio:
            ausencia = "A"
        return {"ausencia": ausencia}
    return {
        "hora_inicio": turno["hora_inicio"],
        "hora_fin": turno["hora_fin"],
    }


def _dias_departamento(
    usuario_id, dias, turnos, versiones, noche_inicio, noche_fin, propio, habituales=None
):
    suyas = versiones.get(usuario_id, [])
    habitual = True if habituales is None else habituales.get(usuario_id, True)
    resultado = []
    for dia in dias:
        turno, _origen = resolver_turno(
            dia,
            turnos.get((usuario_id, dia)),
            suyas,
            noche_inicio,
            noche_fin,
            habitual,
        )
        resultado.append(
            {
                "fecha": dia.isoformat(),
                "turno": _turno_visible(turno, propio),
            }
        )
    return resultado


@bp.get("/turnos/departamento")
@jwt_required()
def listar_turnos_departamento():
    if get_jwt().get("rol") != "trabajador":
        return jsonify(error="No autorizado"), 403
    desde, error = _lunes(request.args.get("desde", ""), False)
    if error:
        return error
    hoy = datetime.now(ZoneInfo("Europe/Madrid")).date()
    lunes_actual = hoy - timedelta(days=hoy.weekday())
    if desde != lunes_actual:
        return jsonify(error="Solo puedes consultar la semana actual"), 403
    dias = _semana(desde)
    yo = get_jwt_identity()

    try:
        with cursor() as cur:
            cur.execute(
                """
                SELECT departamento_id
                FROM usuario
                WHERE id = %s AND rol = 'trabajador'
                """,
                (yo,),
            )
            propio = cur.fetchone()
            if propio is None or propio["departamento_id"] is None:
                return jsonify(error="No autorizado"), 403
            cur.execute(
                """
                SELECT u.id, u.nombre, u.grupo, s.nombre AS seccion
                FROM usuario u
                LEFT JOIN seccion s ON s.id = u.seccion_id
                WHERE u.rol = 'trabajador'
                  AND u.activo
                  AND u.departamento_id = %s
                ORDER BY u.nombre
                """,
                (propio["departamento_id"],),
            )
            trabajadores = cur.fetchall()
            turnos, versiones, noche_inicio, noche_fin, habituales = _material_semana(
                cur, [trabajador["id"] for trabajador in trabajadores], dias
            )
    except SinBaseDeDatos:
        return jsonify(error="Base de datos no configurada"), 503
    except psycopg.Error as exc:
        return error_bd(exc)

    return jsonify(
        desde=desde.isoformat(),
        dias=[dia.isoformat() for dia in dias],
        trabajadores=[
            {
                "id": str(trabajador["id"]),
                "nombre": trabajador["nombre"],
                "grupo": trabajador["grupo"],
                "seccion": trabajador["seccion"],
                "dias": _dias_departamento(
                    trabajador["id"],
                    dias,
                    turnos,
                    versiones,
                    noche_inicio,
                    noche_fin,
                    str(trabajador["id"]) == yo,
                    habituales,
                ),
            }
            for trabajador in trabajadores
        ],
    )


@bp.get("/turnos/mes")
@jwt_required()
def listar_turnos_mes():
    if get_jwt().get("rol") != "mando":
        return jsonify(error="No autorizado"), 403
    dias, error = _mes(request.args.get("anio"), request.args.get("mes"))
    if error:
        return error

    try:
        with cursor() as cur:
            departamento_id, error = _departamento_lectura(cur)
            if error:
                return error
            cur.execute(
                """
                SELECT u.id, u.nombre, u.login, u.grupo, u.activo, s.nombre AS seccion
                FROM usuario u
                LEFT JOIN seccion s ON s.id = u.seccion_id
                WHERE u.rol = 'trabajador' AND u.departamento_id = %s
                ORDER BY u.nombre
                """,
                (departamento_id,),
            )
            trabajadores = cur.fetchall()
            turnos, versiones, noche_inicio, noche_fin, habituales = _material_rango(
                cur, [trabajador["id"] for trabajador in trabajadores], dias[0], dias[-1]
            )
            festivos = _festivos_rango(cur, dias[0], dias[-1])
            trabajadores = [
                trabajador
                for trabajador in trabajadores
                if _visible_en_semana(trabajador, dias, turnos)
            ]
    except SinBaseDeDatos:
        return jsonify(error="Base de datos no configurada"), 503
    except psycopg.Error as exc:
        return error_bd(exc)

    personas = _trabajadores_resueltos(
        trabajadores, dias, turnos, versiones, noche_inicio, noche_fin, festivos, habituales
    )
    for persona, trabajador in zip(personas, trabajadores):
        persona["seccion"] = trabajador["seccion"]
    return jsonify(
        desde=dias[0].isoformat(),
        dias=[dia.isoformat() for dia in dias],
        trabajadores=personas,
    )


@bp.get("/turnos/mios/mes")
@jwt_required()
def listar_turnos_mios_mes():
    if get_jwt().get("rol") != "trabajador":
        return jsonify(error="No autorizado"), 403
    dias, error = _mes(request.args.get("anio"), request.args.get("mes"))
    if error:
        return error

    try:
        with cursor() as cur:
            cur.execute(
                """
                SELECT id, nombre, login, grupo, activo
                FROM usuario
                WHERE id = %s AND rol = 'trabajador'
                """,
                (get_jwt_identity(),),
            )
            trabajador = cur.fetchone()
            if trabajador is None:
                return jsonify(error="No autorizado"), 403
            turnos, versiones, noche_inicio, noche_fin, habituales = _material_rango(
                cur, [trabajador["id"]], dias[0], dias[-1]
            )
            festivos = _festivos_rango(cur, dias[0], dias[-1])
            if not _visible_en_semana(trabajador, dias, turnos):
                trabajador = None
    except SinBaseDeDatos:
        return jsonify(error="Base de datos no configurada"), 503
    except psycopg.Error as exc:
        return error_bd(exc)

    personas = []
    if trabajador is not None:
        personas = _trabajadores_resueltos(
            [trabajador], dias, turnos, versiones, noche_inicio, noche_fin, festivos, habituales
        )
    return jsonify(
        desde=dias[0].isoformat(),
        dias=[dia.isoformat() for dia in dias],
        trabajadores=personas,
    )


def _delegacion_mando(cur, usuario_id):
    """Delegación del departamento al que pertenece el mando."""
    cur.execute(
        """
        SELECT d.delegacion_id
        FROM usuario u
        JOIN departamento d ON d.id = u.departamento_id
        WHERE u.id = %s AND u.rol = 'mando'
        """,
        (usuario_id,),
    )
    fila = cur.fetchone()
    if fila is None or fila["delegacion_id"] is None:
        return None
    return fila["delegacion_id"]


def _personas_mandos(cur, delegacion_id):
    """Mandos activos de esa delegación y quien esté en su departamento mandos."""
    cur.execute(
        """
        SELECT id
        FROM departamento
        WHERE codigo = 'mandos' AND delegacion_id = %s
        """,
        (delegacion_id,),
    )
    mandos = cur.fetchone()
    depto_mandos = mandos["id"] if mandos else None
    cur.execute(
        """
        SELECT u.id, u.nombre, u.grupo, s.nombre AS seccion
        FROM usuario u
        JOIN departamento d ON d.id = u.departamento_id
        LEFT JOIN seccion s ON s.id = u.seccion_id
        WHERE u.activo
          AND d.delegacion_id = %s
          AND (
                u.rol = 'mando'
             OR (%s IS NOT NULL AND u.departamento_id = %s)
          )
        ORDER BY u.nombre
        """,
        (delegacion_id, depto_mandos, depto_mandos),
    )
    return cur.fetchall()


def _cuadrante_mandos(cur, yo, dias):
    delegacion_id = _delegacion_mando(cur, yo)
    if delegacion_id is None:
        return None
    personas = _personas_mandos(cur, delegacion_id)
    turnos, versiones, noche_inicio, noche_fin, habituales = _material_rango(
        cur, [persona["id"] for persona in personas], dias[0], dias[-1]
    )
    return [
        {
            "id": str(persona["id"]),
            "nombre": persona["nombre"],
            "grupo": persona["grupo"],
            "seccion": persona["seccion"],
            "dias": _dias_departamento(
                persona["id"],
                dias,
                turnos,
                versiones,
                noche_inicio,
                noche_fin,
                str(persona["id"]) == str(yo),
                habituales,
            ),
        }
        for persona in personas
    ]


@bp.get("/turnos/mandos")
@jwt_required()
def listar_turnos_mandos():
    if get_jwt().get("rol") != "mando":
        return jsonify(error="No autorizado"), 403
    desde, error = _lunes(request.args.get("desde", ""), False)
    if error:
        return error
    dias = _semana(desde)
    try:
        with cursor() as cur:
            trabajadores = _cuadrante_mandos(cur, get_jwt_identity(), dias)
    except SinBaseDeDatos:
        return jsonify(error="Base de datos no configurada"), 503
    except psycopg.Error as exc:
        return error_bd(exc)
    if trabajadores is None:
        return jsonify(error="No autorizado"), 403
    return jsonify(
        desde=desde.isoformat(),
        dias=[dia.isoformat() for dia in dias],
        trabajadores=trabajadores,
    )


@bp.get("/turnos/mandos/mes")
@jwt_required()
def listar_turnos_mandos_mes():
    if get_jwt().get("rol") != "mando":
        return jsonify(error="No autorizado"), 403
    dias, error = _mes(request.args.get("anio"), request.args.get("mes"))
    if error:
        return error
    try:
        with cursor() as cur:
            trabajadores = _cuadrante_mandos(cur, get_jwt_identity(), dias)
    except SinBaseDeDatos:
        return jsonify(error="Base de datos no configurada"), 503
    except psycopg.Error as exc:
        return error_bd(exc)
    if trabajadores is None:
        return jsonify(error="No autorizado"), 403
    return jsonify(
        desde=dias[0].isoformat(),
        dias=[dia.isoformat() for dia in dias],
        trabajadores=trabajadores,
    )


@bp.put("/turnos/mandos/vacaciones")
@jwt_required()
def guardar_vacaciones_mando():
    if get_jwt().get("rol") != "mando":
        return jsonify(error="No autorizado"), 403
    cuerpo = request.get_json(silent=True) or {}
    yo = get_jwt_identity()
    for clave in ("usuario", "usuario_id", "id"):
        valor = cuerpo.get(clave)
        if valor not in (None, "") and str(valor) != str(yo):
            return jsonify(error="No autorizado"), 403
    ausencia = cuerpo.get("ausencia")
    if ausencia == "":
        ausencia = None
    if ausencia not in (None, "V"):
        return jsonify(error="No autorizado"), 403
    codigo = cuerpo.get("codigo")
    if codigo not in (None, "", "V"):
        return jsonify(error="No autorizado"), 403
    vacaciones = cuerpo.get("vacaciones")
    if not isinstance(vacaciones, bool):
        return jsonify(error="Faltan datos"), 400
    hora_inicio = parse_hora(cuerpo.get("hora_inicio"))
    hora_fin = parse_hora(cuerpo.get("hora_fin"))
    if hora_inicio == "mal" or hora_fin == "mal":
        return jsonify(error="La hora no es válida"), 400
    if hora_inicio is not None or hora_fin is not None:
        return jsonify(error="Indica horario o ausencia, no las dos"), 400
    try:
        fecha = datetime.strptime(cuerpo.get("fecha"), "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return jsonify(error="La fecha no es válida"), 400

    try:
        with cursor() as cur:
            if isinstance(cuerpo.get("login"), str) and cuerpo["login"].strip():
                cur.execute(
                    """
                    SELECT id
                    FROM usuario
                    WHERE id = %s AND login = %s AND rol = 'mando'
                    """,
                    (yo, cuerpo["login"].strip()),
                )
                if cur.fetchone() is None:
                    return jsonify(error="No autorizado"), 403
            if _delegacion_mando(cur, yo) is None:
                return jsonify(error="No autorizado"), 403
            if vacaciones:
                cur.execute(
                    """
                    INSERT INTO turno (
                        usuario_id, fecha, ausencia, hora_inicio, hora_fin,
                        horas_planificadas, horas_nocturnas
                    )
                    VALUES (%s, %s, 'V', NULL, NULL, 0, 0)
                    ON CONFLICT (usuario_id, fecha) DO UPDATE SET
                        ausencia = 'V',
                        hora_inicio = NULL,
                        hora_fin = NULL,
                        horas_planificadas = 0,
                        horas_nocturnas = 0
                    """,
                    (yo, fecha),
                )
            else:
                cur.execute(
                    """
                    DELETE FROM turno
                    WHERE usuario_id = %s AND fecha = %s AND ausencia = 'V'
                    """,
                    (yo, fecha),
                )
    except SinBaseDeDatos:
        return jsonify(error="Base de datos no configurada"), 503
    except psycopg.Error as exc:
        return error_bd(exc)

    return jsonify(fecha=fecha.isoformat(), vacaciones=vacaciones)


@bp.post("/turnos/copiar-semana")
@jwt_required()
def copiar_semana():
    if get_jwt().get("rol") != "mando":
        return jsonify(error="No autorizado"), 403
    cuerpo = request.get_json(silent=True) or {}
    desde, error = _lunes(cuerpo.get("desde", ""), True)
    if error:
        return error
    anterior = desde - timedelta(days=7)
    hasta = desde + timedelta(days=6)

    try:
        with cursor() as cur:
            pedido = cuerpo.get("departamento_id")
            if pedido in (None, ""):
                cur.execute(
                    """
                    SELECT departamento_id
                    FROM usuario
                    WHERE id = %s AND rol = 'mando'
                    """,
                    (get_jwt_identity(),),
                )
                mando = cur.fetchone()
                if mando is None or mando["departamento_id"] is None:
                    return jsonify(error="No autorizado"), 403
                departamento_id = mando["departamento_id"]
            elif not departamento_permitido(cur, get_jwt_identity(), pedido):
                return jsonify(error="No se ha encontrado el departamento"), 404
            else:
                departamento_id = pedido
            cur.execute(
                """
                DELETE FROM turno t
                USING usuario u
                WHERE t.usuario_id = u.id
                  AND u.rol = 'trabajador'
                  AND u.departamento_id = %s
                  AND t.fecha >= %s
                  AND t.fecha <= %s
                  AND NOT EXISTS (
                      SELECT 1
                      FROM turno previo
                      WHERE previo.usuario_id = t.usuario_id
                        AND previo.fecha = t.fecha - 7
                  )
                """,
                (departamento_id, desde, hasta),
            )
            cur.execute(
                """
                INSERT INTO turno (
                    usuario_id, fecha, ausencia, hora_inicio, hora_fin,
                    horas_planificadas, horas_nocturnas
                )
                SELECT t.usuario_id, t.fecha + 7, t.ausencia, t.hora_inicio, t.hora_fin,
                       t.horas_planificadas, t.horas_nocturnas
                FROM turno t
                JOIN usuario u ON u.id = t.usuario_id
                WHERE u.rol = 'trabajador'
                  AND u.departamento_id = %s
                  AND t.fecha >= %s
                  AND t.fecha < %s
                ON CONFLICT (usuario_id, fecha) DO UPDATE SET
                    ausencia = EXCLUDED.ausencia,
                    hora_inicio = EXCLUDED.hora_inicio,
                    hora_fin = EXCLUDED.hora_fin,
                    horas_planificadas = EXCLUDED.horas_planificadas,
                    horas_nocturnas = EXCLUDED.horas_nocturnas
                """,
                (departamento_id, anterior, desde),
            )
            dias_copiados = cur.rowcount
    except SinBaseDeDatos:
        return jsonify(error="Base de datos no configurada"), 503
    except psycopg.Error as exc:
        return error_bd(exc)

    return jsonify(desde=desde.isoformat(), dias_copiados=dias_copiados)


@bp.put("/turnos")
@jwt_required()
def guardar_turno():
    if get_jwt().get("rol") != "mando":
        return jsonify(error="No autorizado"), 403

    cuerpo = request.get_json(silent=True) or {}
    login_nombre = cuerpo.get("login")
    fecha_texto = cuerpo.get("fecha")
    ausencia = cuerpo.get("ausencia")
    if ausencia == "":
        ausencia = None
    hora_inicio = parse_hora(cuerpo.get("hora_inicio"))
    hora_fin = parse_hora(cuerpo.get("hora_fin"))
    if hora_inicio == "mal" or hora_fin == "mal":
        return jsonify(error="La hora no es válida"), 400
    if not isinstance(login_nombre, str) or not login_nombre.strip():
        return jsonify(error="Faltan datos"), 400
    try:
        fecha = datetime.strptime(fecha_texto, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return jsonify(error="La fecha no es válida"), 400

    hay_horario = hora_inicio is not None or hora_fin is not None
    hay_ausencia = ausencia is not None
    if hay_horario and hay_ausencia:
        return jsonify(error="Indica horario o ausencia, no las dos"), 400
    if hay_ausencia:
        if ausencia not in AUSENCIAS:
            return jsonify(error="La ausencia no es válida"), 400
        if hora_inicio is not None or hora_fin is not None:
            return jsonify(error="Indica horario o ausencia, no las dos"), 400
    elif hora_inicio is None or hora_fin is None:
        return jsonify(error="Indica un horario o una ausencia"), 400

    try:
        with cursor() as cur:
            cur.execute(
                """
                SELECT id, departamento_id
                FROM usuario
                WHERE login = %s
                  AND rol = 'trabajador'
                """,
                (login_nombre.strip(),),
            )
            trabajador = cur.fetchone()
            if trabajador is None or not departamento_permitido(
                cur, get_jwt_identity(), trabajador["departamento_id"]
            ):
                return jsonify(error="Ese trabajador no es de tu departamento"), 404
            if hay_ausencia:
                planificadas, nocturnas = 0, 0
                hora_inicio = None
                hora_fin = None
            else:
                cur.execute("SELECT noche_inicio, noche_fin FROM empresa WHERE id = 1")
                franja = cur.fetchone()
                if franja is None:
                    noche_inicio = datetime.strptime("22:00", "%H:%M").time()
                    noche_fin = datetime.strptime("06:00", "%H:%M").time()
                else:
                    noche_inicio = franja["noche_inicio"]
                    noche_fin = franja["noche_fin"]
                planificadas, nocturnas = horas_de_turno(
                    hora_inicio, hora_fin, noche_inicio, noche_fin
                )
            cur.execute(
                """
                INSERT INTO turno (
                    usuario_id, fecha, ausencia, hora_inicio, hora_fin,
                    horas_planificadas, horas_nocturnas
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (usuario_id, fecha) DO UPDATE SET
                    ausencia = EXCLUDED.ausencia,
                    hora_inicio = EXCLUDED.hora_inicio,
                    hora_fin = EXCLUDED.hora_fin,
                    horas_planificadas = EXCLUDED.horas_planificadas,
                    horas_nocturnas = EXCLUDED.horas_nocturnas
                RETURNING fecha, ausencia::text AS ausencia, hora_inicio, hora_fin,
                          horas_planificadas, horas_nocturnas
                """,
                (
                    trabajador["id"],
                    fecha,
                    ausencia,
                    hora_inicio,
                    hora_fin,
                    planificadas,
                    nocturnas,
                ),
            )
            guardado = cur.fetchone()
    except SinBaseDeDatos:
        return jsonify(error="Base de datos no configurada"), 503
    except psycopg.Error as exc:
        return error_bd(exc)

    return jsonify(
        login=login_nombre.strip(),
        fecha=guardado["fecha"].isoformat(),
        ausencia=guardado["ausencia"],
        hora_inicio=hora_texto(guardado["hora_inicio"]),
        hora_fin=hora_texto(guardado["hora_fin"]),
        horas_planificadas=float(guardado["horas_planificadas"]),
        horas_nocturnas=float(guardado["horas_nocturnas"]),
    )


@bp.delete("/turnos")
@jwt_required()
def borrar_turno():
    if get_jwt().get("rol") != "mando":
        return jsonify(error="No autorizado"), 403

    cuerpo = request.get_json(silent=True) or {}
    login_nombre = cuerpo.get("login")
    fecha_texto = cuerpo.get("fecha")
    if not isinstance(login_nombre, str) or not login_nombre.strip():
        return jsonify(error="Faltan datos"), 400
    try:
        fecha = datetime.strptime(fecha_texto, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return jsonify(error="La fecha no es válida"), 400

    try:
        with cursor() as cur:
            cur.execute(
                """
                SELECT id, departamento_id
                FROM usuario
                WHERE login = %s
                  AND rol = 'trabajador'
                """,
                (login_nombre.strip(),),
            )
            trabajador = cur.fetchone()
            if trabajador is None or not departamento_permitido(
                cur, get_jwt_identity(), trabajador["departamento_id"]
            ):
                return jsonify(error="Ese trabajador no es de tu departamento"), 404
            cur.execute(
                "DELETE FROM turno WHERE usuario_id = %s AND fecha = %s",
                (trabajador["id"], fecha),
            )
    except SinBaseDeDatos:
        return jsonify(error="Base de datos no configurada"), 503
    except psycopg.Error as exc:
        return error_bd(exc)

    return jsonify(login=login_nombre.strip(), fecha=fecha.isoformat())


MESES_ARCHIVO = (
    "enero",
    "febrero",
    "marzo",
    "abril",
    "mayo",
    "junio",
    "julio",
    "agosto",
    "septiembre",
    "octubre",
    "noviembre",
    "diciembre",
)
DIAS_LARGOS = ("Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo")
DIAS_CORTOS = ("Lun", "Mar", "Mié", "Jue", "Vie", "Sáb", "Dom")
GRUPOS_EXCEL = (("STEF", "Personal STEF"), ("ETT", "ETT"))


def _hora_compacta(hora):
    horas, minutos = hora.split(":")
    texto = str(int(horas))
    if int(minutos) == 0:
        return texto
    return f"{texto}:{minutos}"


def _texto_celda(turno):
    if not turno:
        return ""
    if turno["ausencia"]:
        return turno["ausencia"]
    return f"{_hora_compacta(turno['hora_inicio'])}-{_hora_compacta(turno['hora_fin'])}"


def _nombre_archivo(codigo, dias, mensual):
    depto = "".join(caracter if caracter.isalnum() else "_" for caracter in codigo).strip("_").upper()
    if not depto:
        depto = "CUADRANTE"
    if mensual:
        return f"{depto}_{MESES_ARCHIVO[dias[0].month - 1]}_{dias[0].year}.xlsx"
    semana = dias[0].isocalendar()
    return f"{depto}_semana_{semana.week}_{semana.year}.xlsx"


def _cabecera_dia(dia, mensual):
    if mensual:
        return f"{DIAS_CORTOS[dia.weekday()]}\n{dia.day}"
    return f"{DIAS_LARGOS[dia.weekday()]}\n{dia.strftime('%d/%m')}"


def _filas_excel(trabajadores):
    filas = []
    vistos = set()
    for grupo, titulo in GRUPOS_EXCEL:
        personas = [trabajador for trabajador in trabajadores if trabajador["grupo"] == grupo]
        if not personas:
            continue
        filas.append(("grupo", titulo))
        for trabajador in personas:
            vistos.add(trabajador["login"])
            filas.append(("persona", trabajador))
    for trabajador in trabajadores:
        if trabajador["login"] not in vistos:
            filas.append(("persona", trabajador))
    return filas


def _fuente_base(libro):
    libro._fonts[0] = Font(name="Arial", size=12)


def _libro(trabajadores, dias, mensual):
    libro = Workbook()
    _fuente_base(libro)
    hoja = libro.active
    hoja.title = "Cuadrante"
    hoja.freeze_panes = "B2"
    fuente = Font(name="Arial", size=12)
    cabecera = Font(name="Arial", size=12, bold=True)
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    grupo_fondo = PatternFill("solid", fgColor="F2F2F2")
    hoja.cell(1, 1, "Trabajador").font = cabecera
    for indice, dia in enumerate(dias, start=2):
        celda = hoja.cell(1, indice, _cabecera_dia(dia, mensual))
        celda.font = cabecera
        celda.alignment = centro
    fila_excel = 2
    for tipo, contenido in _filas_excel(trabajadores):
        if tipo == "grupo":
            celda = hoja.cell(fila_excel, 1, contenido)
            celda.font = cabecera
            celda.fill = grupo_fondo
            if dias:
                hoja.merge_cells(
                    start_row=fila_excel,
                    start_column=1,
                    end_row=fila_excel,
                    end_column=len(dias) + 1,
                )
                for columna in range(2, len(dias) + 2):
                    vacia = hoja.cell(fila_excel, columna)
                    vacia.font = fuente
                    vacia.fill = grupo_fondo
        else:
            hoja.cell(fila_excel, 1, contenido["nombre"]).font = fuente
            for indice, dia in enumerate(contenido["dias"], start=2):
                celda = hoja.cell(fila_excel, indice, _texto_celda(dia["turno"]))
                celda.font = fuente
                celda.alignment = centro
        fila_excel += 1
    hoja.column_dimensions["A"].width = 24
    ancho = 12 if not mensual else 6
    for indice in range(2, len(dias) + 2):
        hoja.column_dimensions[hoja.cell(1, indice).column_letter].width = ancho
    hoja.row_dimensions[1].height = 30
    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return buffer


@bp.get("/turnos/exportar")
@jwt_required()
def exportar_turnos():
    if get_jwt().get("rol") != "mando":
        return jsonify(error="No autorizado"), 403
    anio = request.args.get("anio")
    mes = request.args.get("mes")
    if anio not in (None, "") or mes not in (None, ""):
        dias, error = _mes(anio, mes)
        mensual = True
    elif request.args.get("desde"):
        desde, error = _lunes(request.args.get("desde", ""), False)
        dias = None if error else _semana(desde)
        mensual = False
    else:
        return jsonify(error="Faltan datos"), 400
    if error:
        return error

    try:
        with cursor() as cur:
            departamento_id, error = _departamento_lectura(cur)
            if error:
                return error
            cur.execute(
                "SELECT codigo FROM departamento WHERE id = %s",
                (departamento_id,),
            )
            departamento = cur.fetchone()
            if departamento is None:
                return jsonify(error="No se ha encontrado el departamento"), 404
            cur.execute(
                """
                SELECT id, nombre, login, grupo, activo
                FROM usuario
                WHERE rol = 'trabajador' AND departamento_id = %s
                ORDER BY nombre
                """,
                (departamento_id,),
            )
            trabajadores = cur.fetchall()
            turnos, versiones, noche_inicio, noche_fin, habituales = _material_rango(
                cur, [trabajador["id"] for trabajador in trabajadores], dias[0], dias[-1]
            )
            festivos = _festivos_rango(cur, dias[0], dias[-1])
            trabajadores = [
                trabajador
                for trabajador in trabajadores
                if _visible_en_semana(trabajador, dias, turnos)
            ]
            resueltos = _trabajadores_resueltos(
                trabajadores, dias, turnos, versiones, noche_inicio, noche_fin, festivos, habituales
            )
            codigo = departamento["codigo"]
    except SinBaseDeDatos:
        return jsonify(error="Base de datos no configurada"), 503
    except psycopg.Error as exc:
        return error_bd(exc)

    return send_file(
        _libro(resueltos, dias, mensual),
        as_attachment=True,
        download_name=_nombre_archivo(codigo, dias, mensual),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
