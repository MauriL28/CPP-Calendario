import os
import tempfile
import zipfile
from datetime import datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO
from uuid import UUID

import psycopg
from flask import Blueprint, jsonify, request, send_file
from flask_jwt_extended import get_jwt, get_jwt_identity, jwt_required
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils.exceptions import InvalidFileException

from db import SinBaseDeDatos, cursor
from errores import error_bd
from mandos import departamento_permitido
from fichajes import _segundos, agrupar_dias, cargar, segundos_pausa_con_huecos
from horas import hora_texto
from rutas.turnos import _agrupar_versiones, _franja, _habituales, resolver_turno

bp = Blueprint("fichajes", __name__)

LIMITE_BYTES = 10 * 1024 * 1024
MARGEN_MULTIPARTE = 1024 * 1024
ESTADOS = ("ok", "descuadre", "sin_plan", "sin_emparejar")
ETIQUETAS = {
    "ok": "Correcto",
    "descuadre": "Descuadre",
    "sin_plan": "Sin plan",
    "sin_emparejar": "Sin emparejar",
}
MESES_ARCHIVO = (
    "enero",
    "febrero",
    "marzo",
    "abril",
    "mayo",
    "junio",
    "julio",
    "agosto",
    "septiembre",
    "octubre",
    "noviembre",
    "diciembre",
)


def _nombre(archivo):
    bruto = archivo.filename or ""
    return os.path.basename(bruto.replace("\\", "/")).strip()


@bp.post("/fichajes/importar")
@jwt_required()
def importar():
    if get_jwt().get("rol") != "admin":
        return jsonify(error="No autorizado"), 403

    archivos = [item for item in request.files.values() if item is not None]
    if len(archivos) != 1:
        return jsonify(error="Hay que enviar un solo archivo"), 400
    archivo = archivos[0]
    nombre = _nombre(archivo)
    if not nombre.lower().endswith(".xlsx"):
        return jsonify(error="El archivo tiene que ser .xlsx"), 400

    contenido = archivo.read(LIMITE_BYTES + 1)
    if len(contenido) > LIMITE_BYTES:
        return jsonify(error="El archivo supera 10 MiB"), 400

    temporal = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp.write(contenido)
            temporal = tmp.name
        resultado = cargar(
            temporal,
            registro={
                "usuario_id": get_jwt_identity(),
                "nombre_archivo": nombre,
            },
        )
    except ValueError as exc:
        texto = str(exc).strip() or "El archivo no es un libro Excel válido"
        return jsonify(error=texto), 400
    except (InvalidFileException, zipfile.BadZipFile):
        return jsonify(error="El archivo no es un libro Excel válido"), 400
    except SinBaseDeDatos:
        return jsonify(error="Base de datos no configurada"), 503
    except psycopg.Error as exc:
        return error_bd(exc)
    finally:
        if temporal and os.path.exists(temporal):
            os.remove(temporal)

    return jsonify(
        filas_leidas=resultado["filas_leidas"],
        emparejadas=resultado["filas_con_usuario"],
        sin_emparejar=resultado["sin_emparejar"],
        dias_con_incidencia=resultado["incidencias"],
        filas_nuevas=resultado["filas_nuevas"],
    )


def _fecha(texto):
    if texto is None or texto == "":
        return None
    try:
        return datetime.strptime(texto, "%Y-%m-%d").date()
    except ValueError:
        return False


def _numero(valor):
    if valor is None:
        return None
    return float(valor)


def _argumentos_analisis():
    rol = get_jwt().get("rol")
    if rol not in ("admin", "mando"):
        return None, (jsonify(error="No autorizado"), 403)
    codigos = get_jwt().get("delegaciones")
    if not isinstance(codigos, list) or not codigos:
        return None, (jsonify(error="Delegación no presente"), 401)

    desde = _fecha(request.args.get("desde"))
    hasta = _fecha(request.args.get("hasta"))
    if desde is False or hasta is False:
        return None, (jsonify(error="La fecha no es válida"), 400)
    if desde is not None and hasta is not None and desde > hasta:
        return None, (jsonify(error="desde no puede ser posterior a hasta"), 400)

    estado = request.args.get("coincide") or ""
    if estado and estado not in ESTADOS:
        return None, (jsonify(error="coincide no es válido"), 400)

    departamento = request.args.get("departamento_id") or ""
    if rol == "admin" and departamento:
        try:
            UUID(departamento)
        except ValueError:
            return None, (jsonify(error="departamento_id no es válido"), 400)

    return {
        "rol": rol,
        "codigos": codigos,
        "desde": desde,
        "hasta": hasta,
        "estado": estado,
        "departamento": departamento,
        "identidad": get_jwt_identity(),
    }, None


def _alcance_analisis(cur, ctx):
    departamento = ctx["departamento"]
    if ctx["rol"] == "mando":
        if departamento:
            if not departamento_permitido(cur, ctx["identidad"], departamento):
                return None, None, (jsonify(error="No se ha encontrado el departamento"), 404)
        else:
            cur.execute(
                "SELECT departamento_id FROM usuario WHERE id = %s",
                (ctx["identidad"],),
            )
            propio = cur.fetchone()
            if propio is None or propio["departamento_id"] is None:
                return None, None, (jsonify(error="No autorizado"), 403)
            departamento = str(propio["departamento_id"])
    elif departamento:
        cur.execute(
            """
            SELECT d.id
            FROM departamento d
            JOIN delegacion g ON g.id = d.delegacion_id
            WHERE d.id = %s
              AND g.codigo = ANY(%s)
            """,
            (departamento, ctx["codigos"]),
        )
        if cur.fetchone() is None:
            return None, None, (jsonify(error="No autorizado"), 403)

    clausulas = [
        """(
            f.coincide IS NOT NULL
            OR (
                f.incidencia
                AND EXISTS (
                    SELECT 1
                    FROM fichaje AS otra
                    WHERE otra.fecha = f.fecha
                      AND otra.id <> f.id
                      AND (
                            (f.usuario_id IS NOT NULL AND otra.usuario_id = f.usuario_id)
                         OR (f.usuario_id IS NULL AND otra.numero_sap = f.numero_sap)
                      )
                )
            )
        )"""
    ]
    params = []
    if ctx["rol"] == "mando" or departamento:
        clausulas.append("u.departamento_id = %s")
        params.append(departamento)
    else:
        clausulas.append("(f.usuario_id IS NULL OR g.codigo = ANY(%s))")
        params.append(ctx["codigos"])
    if ctx["estado"]:
        clausulas.append("f.coincide = %s")
        params.append(ctx["estado"])
    if ctx["desde"] is not None:
        clausulas.append("f.fecha >= %s")
        params.append(ctx["desde"])
    if ctx["hasta"] is not None:
        clausulas.append("f.fecha <= %s")
        params.append(ctx["hasta"])
    return " AND ".join(clausulas), params, None


def _minutos(segundos):
    if segundos is None:
        return None
    return int((Decimal(segundos) / Decimal(60)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _minutos_reloj(inicio, fin):
    if inicio is None or fin is None:
        return None
    delta = _segundos(inicio, fin)
    if delta < 0:
        delta += Decimal(24 * 3600)
    return _minutos(delta)


def _hhmm(valor):
    if valor is None:
        return None
    if isinstance(valor, str):
        return valor[:5]
    return hora_texto(valor)


def _eventos_analisis(cur, filas):
    """Jornadas cuya fecha cae en las filas. Incluye el día anterior y dos después."""
    fechas = [fila["fecha"] for fila in filas]
    if not fechas:
        return {}
    ids = list({fila["usuario_id"] for fila in filas if fila["usuario_id"] is not None})
    saps = list({fila["numero_sap"] for fila in filas if fila["numero_sap"]})
    cur.execute(
        """
        SELECT MIN(fecha + hora) AS inicio, MAX(fecha + hora) AS fin
        FROM fichaje_evento
        WHERE fecha IS NOT NULL
        """
    )
    extremo = cur.fetchone()
    rango = None
    if extremo and extremo["inicio"] is not None and extremo["fin"] is not None:
        rango = (extremo["inicio"], extremo["fin"])
    cur.execute(
        """
        SELECT numero_sap, usuario_id, tipo, fecha, hora, orden
        FROM fichaje_evento
        WHERE fecha BETWEEN %s AND %s
          AND (usuario_id = ANY(%s) OR numero_sap = ANY(%s))
        ORDER BY fecha, hora, orden
        """,
        (min(fechas) - timedelta(days=1), max(fechas) + timedelta(days=2), ids, saps),
    )
    grupos = {}
    for dia in agrupar_dias(cur.fetchall(), rango):
        clave = (dia["usuario_id"] or dia["numero_sap"], dia["fecha"])
        grupos.setdefault(clave, []).append(dia)
    return grupos


def _tomar_jornada(restantes, fila):
    """Cada fila de fichaje se queda con su jornada, también si hay incidencia."""
    if not restantes:
        return None
    mismas = [dia for dia in restantes if bool(dia["incidencia"]) == bool(fila["incidencia"])]
    candidatas = mismas or list(restantes)
    if not fila["incidencia"]:
        reales = [dia for dia in candidatas if dia.get("motivo") != "fuera de rango"]
        if reales:
            candidatas = reales
    elegido = candidatas[0]
    restantes.remove(elegido)
    return elegido


def _planes_analisis(cur, filas):
    utiles = [fila for fila in filas if fila["usuario_id"] is not None]
    if not utiles:
        return {}
    ids = list({fila["usuario_id"] for fila in utiles})
    fechas = [fila["fecha"] for fila in utiles]
    cur.execute(
        """
        SELECT usuario_id, fecha, ausencia::text AS ausencia,
               hora_inicio, hora_fin, horas_planificadas, horas_nocturnas
        FROM turno
        WHERE usuario_id = ANY(%s)
          AND fecha = ANY(%s)
        """,
        (ids, fechas),
    )
    turnos = {(fila["usuario_id"], fila["fecha"]): fila for fila in cur.fetchall()}
    cur.execute(
        """
        SELECT v.usuario_id, v.desde, d.dia_semana, d.horario_inicio, d.horario_fin
        FROM horario_version v
        LEFT JOIN horario_dia d
          ON d.usuario_id = v.usuario_id AND d.desde = v.desde
        WHERE v.usuario_id = ANY(%s)
          AND v.desde <= %s
        """,
        (ids, max(fechas)),
    )
    versiones = _agrupar_versiones(cur.fetchall())
    noche_inicio, noche_fin = _franja(cur)
    habituales = _habituales(cur, ids)
    planes = {}
    for fila in utiles:
        guardado = turnos.get((fila["usuario_id"], fila["fecha"]))
        turno, _origen = resolver_turno(
            fila["fecha"],
            guardado,
            versiones.get(fila["usuario_id"], []),
            noche_inicio,
            noche_fin,
            habituales.get(fila["usuario_id"], True),
        )
        planes[(fila["usuario_id"], fila["fecha"])] = turno
    return planes


def _detalle_fila(fila, dia, turno):
    lista = dia["eventos"] if dia else []
    llegadas = [evento for evento in lista if evento["tipo"] == "Llegada"]
    salidas = [evento for evento in lista if evento["tipo"] == "Salida"]
    entrada = _hhmm(llegadas[0]["hora"]) if llegadas else None
    salida_ev = salidas[-1] if salidas else None
    salida = _hhmm(salida_ev["hora"]) if salida_ev else None
    offset = 0
    if salida_ev is not None and salida_ev.get("fecha") is not None and fila["fecha"] is not None:
        offset = 1 if (salida_ev["fecha"] - fila["fecha"]).days >= 1 else 0
    fichado = _minutos(dia["segundos_trabajo"]) if dia else None
    pausa = _minutos(segundos_pausa_con_huecos(lista)) if lista else None
    inicio = turno.get("hora_inicio") if turno else None
    fin = turno.get("hora_fin") if turno else None
    previsto = _minutos_reloj(
        datetime.strptime(inicio, "%H:%M").time() if isinstance(inicio, str) else inicio,
        datetime.strptime(fin, "%H:%M").time() if isinstance(fin, str) else fin,
    ) if inicio and fin else None
    diferencia = None if previsto is None or fichado is None else fichado - previsto
    return {
        "previsto_min": previsto,
        "fichado_min": fichado,
        "pausa_min": pausa,
        "diferencia_min": diferencia,
        "entrada_fichada": entrada,
        "salida_fichada": salida,
        "salida_dia_offset": offset,
        "entrada_prevista": _hhmm(inicio) if inicio else None,
        "salida_prevista": _hhmm(fin) if fin else None,
    }


def _filas_analisis(cur, donde, params):
    cur.execute(
        f"""
        SELECT
            f.usuario_id,
            f.numero_sap,
            CASE
                WHEN f.usuario_id IS NULL THEN f.numero_sap
                ELSE u.nombre
            END AS nombre,
            CASE
                WHEN f.usuario_id IS NULL THEN NULL
                ELSE d.nombre
            END AS departamento,
            f.fecha,
            f.incidencia,
            f.horas_planificadas,
            f.horas_trabajadas,
            f.horas_pausa,
            CASE
                WHEN f.horas_planificadas IS NULL THEN NULL
                ELSE f.horas_trabajadas - f.horas_planificadas
            END AS diferencia,
            f.coincide::text AS coincide
        FROM fichaje f
        LEFT JOIN usuario u ON u.id = f.usuario_id
        LEFT JOIN departamento d ON d.id = u.departamento_id
        LEFT JOIN delegacion g ON g.id = d.delegacion_id
        WHERE {donde}
        ORDER BY f.fecha, nombre
        """,
        params,
    )
    leidas = cur.fetchall()
    eventos = _eventos_analisis(cur, leidas)
    planes = _planes_analisis(cur, leidas)
    pendientes = {clave: list(dias) for clave, dias in eventos.items()}
    filas = []
    for fila in leidas:
        clave = (fila["usuario_id"] or fila["numero_sap"], fila["fecha"])
        detalle = _detalle_fila(
            fila,
            _tomar_jornada(pendientes.setdefault(clave, []), fila),
            planes.get((fila["usuario_id"], fila["fecha"])),
        )
        filas.append(
            {
                "nombre": fila["nombre"],
                "departamento": fila["departamento"],
                "fecha": fila["fecha"].isoformat(),
                "horas_planificadas": _numero(fila["horas_planificadas"]),
                "horas_trabajadas": _numero(fila["horas_trabajadas"]),
                "horas_pausa": _numero(fila["horas_pausa"]),
                "diferencia": _numero(fila["diferencia"]),
                "coincide": fila["coincide"],
                **detalle,
            }
        )
    return filas


def _horas_texto(valor):
    if valor is None:
        return "—"
    return f"{valor:.2f}".replace(".", ",")


def _diferencia_texto(valor):
    if valor is None:
        return "—"
    texto = f"{abs(valor):.1f}".replace(".", ",")
    if valor > 0:
        return f"+{texto}"
    if valor < 0:
        return f"-{texto}"
    return texto


def _nombre_analisis(estado, desde, hasta):
    partes = ["analisis"]
    if estado:
        partes.append(estado)
    if desde is not None and hasta is not None and desde.year == hasta.year and desde.month == hasta.month:
        partes.append(MESES_ARCHIVO[desde.month - 1])
        partes.append(str(desde.year))
    elif desde is not None and hasta is not None and desde.year == hasta.year:
        partes.append(str(desde.year))
    elif desde is not None and hasta is not None:
        partes.append(desde.isoformat())
        partes.append(hasta.isoformat())
    else:
        anio = desde.year if desde is not None else hasta.year if hasta is not None else datetime.now().year
        partes.append(str(anio))
    return "_".join(partes) + ".xlsx"


def _libro_analisis(filas):
    libro = Workbook()
    libro._fonts[0] = Font(name="Arial", size=12)
    hoja = libro.active
    hoja.title = "Análisis"
    fuente = Font(name="Arial", size=12)
    cabecera = Font(name="Arial", size=12, bold=True)
    centro = Alignment(horizontal="center")
    columnas = (
        "Trabajador",
        "Departamento",
        "Fecha",
        "Previsto",
        "Fichado",
        "Pausa",
        "Diferencia",
        "Entrada prevista",
        "Salida prevista",
        "Entrada fichada",
        "Salida fichada",
        "Previsto (min)",
        "Fichado (min)",
        "Pausa (min)",
        "Diferencia (min)",
        "Estado",
    )
    for indice, titulo in enumerate(columnas, start=1):
        celda = hoja.cell(1, indice, titulo)
        celda.font = cabecera
    for fila_excel, fila in enumerate(filas, start=2):
        hoja.cell(fila_excel, 1, fila["nombre"] or "").font = fuente
        hoja.cell(fila_excel, 2, fila["departamento"] or "").font = fuente
        hoja.cell(
            fila_excel,
            3,
            datetime.strptime(fila["fecha"], "%Y-%m-%d").strftime("%d/%m/%Y"),
        ).font = fuente
        for columna, texto in enumerate(
            (
                _horas_texto(fila["horas_planificadas"]),
                _horas_texto(fila["horas_trabajadas"]),
                _horas_texto(None if fila["pausa_min"] is None else fila["pausa_min"] / 60),
                _diferencia_texto(fila["diferencia"]),
            ),
            start=4,
        ):
            celda = hoja.cell(fila_excel, columna, texto)
            celda.font = fuente
            celda.alignment = centro
        for columna, texto in enumerate(
            (
                fila["entrada_prevista"] or "—",
                fila["salida_prevista"] or "—",
                fila["entrada_fichada"] or "—",
                fila["salida_fichada"] or "—",
            ),
            start=8,
        ):
            celda = hoja.cell(fila_excel, columna, texto)
            celda.font = fuente
            celda.alignment = centro
        for columna, minutos in enumerate(
            (
                fila["previsto_min"],
                fila["fichado_min"],
                fila["pausa_min"],
                fila["diferencia_min"],
            ),
            start=12,
        ):
            celda = hoja.cell(fila_excel, columna, minutos)
            celda.font = fuente
            celda.alignment = centro
            celda.number_format = "0"
        hoja.cell(fila_excel, 16, ETIQUETAS.get(fila["coincide"], fila["coincide"])).font = fuente
    anchos = (24, 22, 12, 12, 12, 12, 12, 16, 16, 16, 16, 14, 14, 12, 16, 16)
    for indice, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[hoja.cell(1, indice).column_letter].width = ancho
    hoja.freeze_panes = "A2"
    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return buffer


@bp.get("/fichajes/analisis")
@jwt_required()
def analisis():
    ctx, error = _argumentos_analisis()
    if error:
        return error

    try:
        with cursor() as cur:
            donde, params, error = _alcance_analisis(cur, ctx)
            if error:
                return error
            filas = _filas_analisis(cur, donde, params)
            cur.execute(
                f"""
                SELECT
                    COUNT(*) AS total,
                    COUNT(*) FILTER (WHERE f.coincide = 'ok') AS ok,
                    COUNT(*) FILTER (WHERE f.coincide = 'descuadre') AS descuadre,
                    COUNT(*) FILTER (WHERE f.coincide = 'sin_plan') AS sin_plan,
                    COUNT(*) FILTER (WHERE f.coincide = 'sin_emparejar') AS sin_emparejar
                FROM fichaje f
                LEFT JOIN usuario u ON u.id = f.usuario_id
                LEFT JOIN departamento d ON d.id = u.departamento_id
                LEFT JOIN delegacion g ON g.id = d.delegacion_id
                WHERE {donde}
                """,
                params,
            )
            totales = cur.fetchone()
    except SinBaseDeDatos:
        return jsonify(error="Base de datos no configurada"), 503
    except psycopg.Error as exc:
        return error_bd(exc)

    return jsonify(
        filas=filas,
        totales={
            "total": totales["total"],
            "ok": totales["ok"],
            "descuadre": totales["descuadre"],
            "sin_plan": totales["sin_plan"],
            "sin_emparejar": totales["sin_emparejar"],
        },
    )


@bp.get("/fichajes/analisis/exportar")
@jwt_required()
def exportar_analisis():
    ctx, error = _argumentos_analisis()
    if error:
        return error

    try:
        with cursor() as cur:
            donde, params, error = _alcance_analisis(cur, ctx)
            if error:
                return error
            filas = _filas_analisis(cur, donde, params)
    except SinBaseDeDatos:
        return jsonify(error="Base de datos no configurada"), 503
    except psycopg.Error as exc:
        return error_bd(exc)

    return send_file(
        _libro_analisis(filas),
        as_attachment=True,
        download_name=_nombre_analisis(ctx["estado"], ctx["desde"], ctx["hasta"]),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
